import openpyxl
from datetime import datetime

# File path
file_path = "data.xlsx"

# Button states (simulated, replace with your logic)
current_state = "On"  # Example: dynamically set based on the HTML
last_state = None

try:
    # Load workbook and sheet
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active

    # Get the last logged state (if any)
    if sheet.max_row > 1:
        last_state = sheet.cell(row=sheet.max_row, column=2).value

    # Log only if the state changes
    if current_state != last_state:
        now = datetime.now().strftime("%I:%M %p")  # Time in 12-hour format
        sheet.append([now, current_state])
        wb.save(file_path)
        print(f"Logged: {current_state} at {now}")
    else:
        print("No state change. No log added.")

except FileNotFoundError:
    # Create new file if it doesn't exist
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Time", "State"])
    now = datetime.now().strftime("%I:%M %p")
    sheet.append([now, current_state])
    wb.save(file_path)
    print(f"Created new log: {current_state} at {now}")
